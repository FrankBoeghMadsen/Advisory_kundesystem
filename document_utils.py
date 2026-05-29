
from pathlib import Path

def extract_text_from_file(path: str) -> str:
    p = Path(path)
    suffix = p.suffix.lower()
    try:
        if suffix == ".txt":
            return p.read_text(encoding="utf-8", errors="ignore")
        if suffix == ".docx":
            from docx import Document
            doc = Document(str(p))
            parts = []
            for para in doc.paragraphs:
                if para.text.strip():
                    parts.append(para.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            return "\n".join(parts)
        if suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(p))
            parts = []
            for page in reader.pages:
                txt = page.extract_text() or ""
                if txt.strip():
                    parts.append(txt.strip())
            return "\n\n".join(parts)
    except Exception as e:
        return f"[Kunne ikke udtrække tekst: {e}]"
    return "[Filtypen understøttes ikke til tekstudtræk endnu.]"

def simple_meeting_summary(text: str) -> dict:
    text = (text or "").strip()
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    short = " ".join(lines[:8])[:1200]
    themes = []
    for kw in ["GMP", "GDP", "QP", "RP", "QA", "RA", "LMST", "Lægemiddelstyrelsen", "recall", "tilbagekald", "inspektion", "myndighed", "compliance", "ledelse", "grossist", "fremstilling"]:
        if kw.lower() in text.lower():
            themes.append(kw)
    return {
        "summary": short,
        "key_takeaways": "Mulige temaer: " + ", ".join(dict.fromkeys(themes)) if themes else "Ingen tydelige nøgletemaer fundet automatisk.",
        "next_steps": "Gennemgå referatet manuelt og opret relevante observationer eller opfølgningspunkter."
    }
